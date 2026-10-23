"""Worksheet-backed NFT number lookup/editing, independent of the Qt UI."""
import csv
from pathlib import Path
from openpyxl import load_workbook, Workbook

class NFTWorksheet:
    def __init__(self):
        self.headers=[]
        self.rows=[]
        self.source=None
        self.sheet_name=None
        self.number_column=None

    @staticmethod
    def normalize_number(value):
        text=str(value).strip() if value is not None else ''
        if not text:return ''
        # Numeric worksheet cells often drop leading zeros; treat 0059 and 59 as equivalent.
        try:
            from decimal import Decimal
            n=Decimal(text)
            if n>=0 and n==int(n):return str(int(n))
        except Exception:pass
        return text.casefold()

    def load(self,path,sheet_name=None):
        path=Path(path)
        if path.suffix.lower()=='.csv':
            with path.open('r',encoding='utf-8-sig',newline='') as f:
                records=list(csv.reader(f))
            name=None
        elif path.suffix.lower()=='.xlsx':
            wb=load_workbook(path,read_only=True,data_only=True)
            try:
                if sheet_name and sheet_name not in wb.sheetnames:raise ValueError('Worksheet tab not found')
                ws=wb[sheet_name] if sheet_name else wb.active
                records=list(ws.values)
                name=ws.title
            finally:wb.close()
        else:raise ValueError('Choose an .xlsx or .csv worksheet')
        if not records:raise ValueError('Worksheet is empty')
        headers=[str(c).strip() if c is not None else '' for c in records[0]]
        if not any(headers):raise ValueError('First row must contain column names')
        if len(set(headers))!=len(headers) or any(not h for h in headers):raise ValueError('Column names must be unique and nonempty')
        self.headers=headers
        self.rows=[{h:('' if i>=len(row) or row[i] is None else str(row[i])) for i,h in enumerate(headers)} for row in records[1:] if any(c is not None and str(c).strip() for c in row)]
        self.source=str(path);self.sheet_name=name
        preferred=('nft number','nft #','nft','number','token id','token_id','id','collection number','collection #')
        self.number_column=next((h for h in headers if h.strip().casefold() in preferred),headers[0])
        return len(self.rows)

    def lookup(self,number):
        if not self.number_column:raise ValueError('Load a worksheet first')
        key=self.normalize_number(number)
        if not key:raise ValueError('Enter an NFT number')
        matches=[i for i,row in enumerate(self.rows) if self.normalize_number(row[self.number_column])==key]
        if not matches:raise LookupError(f'NFT {number} not found in column {self.number_column}')
        if len(matches)>1:raise ValueError(f'NFT number {number} appears {len(matches)} times; make the ID column unique')
        return matches[0],self.rows[matches[0]].copy()

    def edit(self,index,values):
        if not 0<=index<len(self.rows):raise IndexError('NFT row out of range')
        if any(key not in self.headers for key in values):raise ValueError('Unknown worksheet column')
        self.rows[index].update({k:str(v) for k,v in values.items()})

    def export(self,path):
        path=Path(path)
        if not self.headers:raise ValueError('Load a worksheet first')
        if path.suffix.lower()=='.csv':
            with path.open('w',encoding='utf-8-sig',newline='') as f:
                w=csv.DictWriter(f,fieldnames=self.headers);w.writeheader();w.writerows(self.rows)
        elif path.suffix.lower()=='.xlsx':
            wb=Workbook();ws=wb.active;ws.title=(self.sheet_name or 'NFT Traits')[:31]
            ws.append(self.headers)
            for row in self.rows:ws.append([row.get(h,'') for h in self.headers])
            wb.save(path)
        else:raise ValueError('Export must be .xlsx or .csv')
        return path
